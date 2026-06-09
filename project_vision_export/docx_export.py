from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .branding import COLORS, LOGO_PATH
from .formatting import format_project_vision_data


def _rgb(value):
    value = value.lstrip("#")
    return RGBColor.from_string(value)


def _set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill.lstrip("#"))


def _set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_cell_border(cell, **edges):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge_name, settings in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in settings.items():
            edge.set(qn(f"w:{key}"), str(value))


def _set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_fixed_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_width_dxa = sum(int(width.inches * 1440) for width in widths)
    properties = table._tbl.tblPr
    width_element = properties.first_child_found_in("w:tblW")
    width_element.set(qn("w:w"), str(table_width_dxa))
    width_element.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(width.inches * 1440)))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(int(widths[index].inches * 1440)))
            tc_width.set(qn("w:type"), "dxa")


def _set_run(run, size=9, bold=False, color=None, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)


def _clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _add_text(paragraph, text, **formatting):
    run = paragraph.add_run(text)
    _set_run(run, **formatting)
    return run


def _configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (
        ("Heading 1", 11.5, 8, 5),
        ("Heading 2", 10.5, 7, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(COLORS["red"])
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _enable_field_updates(document):
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _add_letterhead(header, data):
    table = header.add_table(rows=2, cols=3, width=Cm(17.4))
    widths = [Cm(4.4), Cm(8.4), Cm(4.6)]
    _set_fixed_table_width(table, widths)

    logo_cell, details_cell, status_cell = table.rows[0].cells
    logo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    logo_paragraph = logo_cell.paragraphs[0]
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.add_run().add_picture(str(LOGO_PATH), width=Cm(4.1))

    details = details_cell.paragraphs[0]
    details.paragraph_format.space_after = Pt(0)
    _add_text(details, data["company"]["name"], size=14, bold=True, color=COLORS["black"])
    for line in (
        data["company"]["address"],
        data["company"]["phone"],
        data["company"]["email"],
        f"Reg: {data['company']['registration']}",
    ):
        details.add_run().add_break()
        _add_text(details, line, size=7.8, color=COLORS["slate"])

    status = status_cell.paragraphs[0]
    status.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_text(status, "PROJECT PLANNING", size=11.5, bold=True, color=COLORS["red"])
    status.add_run().add_break()
    _add_text(status, data["reference"], size=7.5, bold=True, color=COLORS["black"])
    status.add_run().add_break()
    status.add_run().add_break()
    _add_text(status, "PRELIMINARY", size=7.5, bold=True, color=COLORS["red"])

    tagline = table.rows[1].cells[0]
    tagline.merge(table.rows[1].cells[2])
    _set_cell_shading(tagline, COLORS["red"])
    tag_p = tagline.paragraphs[0]
    tag_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag_p.paragraph_format.space_before = Pt(2)
    tag_p.paragraph_format.space_after = Pt(2)
    _add_text(tag_p, data["company"]["tagline"], size=8, bold=True, color=COLORS["white"])

    for row in table.rows:
        for cell in row.cells:
            _set_cell_margins(cell, top=40, start=60, bottom=40, end=60)
            _set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                start={"val": "nil"},
                end={"val": "nil"},
            )


def _add_footer(footer, data):
    line = footer.add_paragraph()
    line.paragraph_format.space_after = Pt(2)
    properties = line._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:top")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), COLORS["line"].lstrip("#"))
    border.append(bottom)
    properties.append(border)

    table = footer.add_table(rows=1, cols=2, width=Cm(17.4))
    _set_fixed_table_width(table, [Cm(14.5), Cm(2.9)])
    left, right = table.rows[0].cells
    left_p = left.paragraphs[0]
    _add_text(
        left_p,
        f"{data['company']['name']} | {data['company']['phone']} | {data['company']['email']}",
        size=7,
        color=COLORS["slate"],
    )
    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_text(right_p, "Page ", size=7, color=COLORS["slate"])
    _add_field(right_p, "PAGE")
    _add_text(right_p, " of ", size=7, color=COLORS["slate"])
    _add_field(right_p, "NUMPAGES")
    for cell in (left, right):
        _set_cell_border(
            cell,
            top={"val": "nil"},
            bottom={"val": "nil"},
            start={"val": "nil"},
            end={"val": "nil"},
        )


def _add_field(paragraph, instruction):
    run = paragraph.add_run()
    _set_run(run, size=7, color=COLORS["slate"])
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, display, end])


def _add_label_value_table(document, rows):
    table = document.add_table(rows=0, cols=2)
    for label, value in rows:
        cells = table.add_row().cells
        label_p = cells[0].paragraphs[0]
        value_p = cells[1].paragraphs[0]
        _add_text(label_p, label, size=8.2, bold=True, color=COLORS["slate"])
        _add_text(value_p, str(value), size=8.6, color=COLORS["black"])
        for cell in cells:
            _set_cell_shading(cell, COLORS["panel"])
            _set_cell_margins(cell, top=80, start=110, bottom=80, end=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": COLORS["line"].lstrip("#")},
                bottom={"val": "single", "sz": "4", "color": COLORS["line"].lstrip("#")},
                start={"val": "single", "sz": "4", "color": COLORS["line"].lstrip("#")},
                end={"val": "single", "sz": "4", "color": COLORS["line"].lstrip("#")},
            )
    _set_fixed_table_width(table, [Cm(4.2), Cm(13.2)])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _add_section_heading(document, text):
    document.add_paragraph(text, style="Heading 1")


def _add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.3)
    paragraph.paragraph_format.space_after = Pt(3)
    _clear_paragraph(paragraph)
    _add_text(paragraph, text, size=9.1, color=COLORS["black"])


def _add_callout(document, title, body, fill=None):
    table = document.add_table(rows=1, cols=2)
    _set_fixed_table_width(table, [Cm(4.5), Cm(12.9)])
    left, right = table.rows[0].cells
    for cell in (left, right):
        _set_cell_shading(cell, fill or COLORS["soft_red"])
        _set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
        _set_cell_border(
            cell,
            top={"val": "single", "sz": "6", "color": COLORS["red"].lstrip("#")},
            bottom={"val": "single", "sz": "6", "color": COLORS["red"].lstrip("#")},
            start={"val": "single", "sz": "6", "color": COLORS["red"].lstrip("#")},
            end={"val": "single", "sz": "6", "color": COLORS["red"].lstrip("#")},
        )
    _add_text(left.paragraphs[0], title, size=9.5, bold=True, color=COLORS["red"])
    _add_text(right.paragraphs[0], body, size=9.1, color=COLORS["black"])


def _add_readiness(document, data):
    _add_section_heading(document, "Project Readiness Overview")
    score = data["readiness_score"]
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    _add_text(
        paragraph,
        f"Project Readiness Score: {score} / 100",
        size=11.5,
        bold=True,
        color=COLORS["black"],
    )

    table = document.add_table(rows=1, cols=2)
    filled_units = max(score, 1)
    empty_units = max(100 - score, 1)
    total_units = filled_units + empty_units
    filled = filled_units / total_units
    total_width = 17.4
    _set_fixed_table_width(
        table,
        [Cm(total_width * filled), Cm(total_width * (1 - filled))],
    )
    _set_cell_shading(table.cell(0, 0), COLORS["red"])
    _set_cell_shading(table.cell(0, 1), "#E4E7EC")
    for cell in table.rows[0].cells:
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        _set_cell_margins(cell, top=50, start=0, bottom=50, end=0)
        _set_cell_border(
            cell,
            top={"val": "nil"},
            bottom={"val": "nil"},
            start={"val": "nil"},
            end={"val": "nil"},
        )
    explanation = document.add_paragraph()
    explanation.paragraph_format.space_before = Pt(5)
    _add_text(
        explanation,
        f"{data['readiness_label']}. ",
        size=9.1,
        bold=True,
        color=COLORS["black"],
    )
    _add_text(explanation, data["readiness_explanation"], size=9.1, color=COLORS["black"])


def generate_project_vision_docx(
    project_data,
    company_settings=None,
    output_path=None,
    reference=None,
):
    data = format_project_vision_data(
        project_data,
        company_settings=company_settings,
        reference=reference,
    )
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(5.8)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.55)
    _configure_styles(document)
    _enable_field_updates(document)
    _add_letterhead(section.header, data)
    _add_footer(section.footer, data)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(1)
    _add_text(title, "PROJECT VISION SUMMARY", size=20, bold=True, color=COLORS["red"])
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    _add_text(
        subtitle,
        data["brief_subtitle"],
        size=10,
        color=COLORS["slate"],
    )

    _add_label_value_table(
        document,
        [
            ("Document Reference", data["reference"]),
            ("Date Generated", data["generated_date"]),
            ("Prepared For", data["client_name"] or "Prospective Client"),
            ("Project Category", data["project_category"]),
            ("Prepared By", data["prepared_by"]),
            ("Document Status", data["document_status"]),
        ],
    )
    intro = document.add_paragraph()
    intro.paragraph_format.space_before = Pt(2)
    _add_text(intro, data["introduction"], size=9.2, color=COLORS["black"])

    _add_section_heading(document, "Client Vision Overview")
    _add_label_value_table(
        document,
        [
            ("Project Type", data["project_type"]),
            ("Preferred Style", data["preferred_style"]),
            ("Lifestyle Goal", data["lifestyle_goal"]),
            ("Current Project Stage", data["current_stage"]),
            ("Intended Timeline", data["timeline"]),
            ("Budget Direction", data["budget"]),
            ("Project Location", data["project_location"]),
            ("Plot Size", data["plot_size"]),
            ("Target Floor Area", data["floor_area"]),
            ("Bedrooms / Main Rooms", data["bedrooms"]),
            ("Bathrooms", data["bathrooms"]),
            ("Building Levels", data["storeys"]),
        ],
    )
    feature_heading = document.add_paragraph()
    feature_heading.paragraph_format.space_before = Pt(3)
    feature_heading.paragraph_format.space_after = Pt(2)
    _add_text(feature_heading, "Selected Features", size=9, bold=True, color=COLORS["slate"])
    for feature in data["features"]:
        _add_bullet(document, feature)

    _add_callout(
        document,
        "Recommended Design Direction",
        data["design_direction"],
    )
    _add_section_heading(document, "Recommended Professional Services")
    for service in data["services"]:
        _add_bullet(document, service)

    _add_readiness(document, data)
    _add_callout(
        document,
        "Recommended Next Step",
        data["next_step"],
        fill=COLORS["panel"],
    )
    contact = document.add_table(rows=1, cols=2)
    _set_fixed_table_width(contact, [Cm(5.2), Cm(12.2)])
    left, right = contact.rows[0].cells
    for cell in (left, right):
        _set_cell_shading(cell, COLORS["soft_red"])
        _set_cell_margins(cell, top=100, start=130, bottom=100, end=130)
        _set_cell_border(
            cell,
            top={"val": "single", "sz": "5", "color": COLORS["red"].lstrip("#")},
            bottom={"val": "single", "sz": "5", "color": COLORS["red"].lstrip("#")},
            start={"val": "single", "sz": "5", "color": COLORS["red"].lstrip("#")},
            end={"val": "single", "sz": "5", "color": COLORS["red"].lstrip("#")},
        )
    _add_text(
        left.paragraphs[0],
        "BOOK A PROFESSIONAL CONSULTATION",
        size=8.3,
        bold=True,
        color=COLORS["red"],
    )
    contact_p = right.paragraphs[0]
    _add_text(contact_p, "WhatsApp: ", size=8.5, bold=True, color=COLORS["slate"])
    _add_text(contact_p, data["company"]["whatsapp"], size=8.5, color=COLORS["black"])
    contact_p.add_run().add_break()
    _add_text(contact_p, "Phone: ", size=8.5, bold=True, color=COLORS["slate"])
    _add_text(contact_p, data["company"]["phone"], size=8.5, color=COLORS["black"])
    contact_p.add_run().add_break()
    _add_text(contact_p, "Email: ", size=8.5, bold=True, color=COLORS["slate"])
    _add_text(contact_p, data["company"]["email"], size=8.5, color=COLORS["black"])

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(7)
    disclaimer.paragraph_format.space_after = Pt(0)
    _add_text(disclaimer, "Disclaimer: ", size=7.6, bold=True, color=COLORS["slate"])
    _add_text(disclaimer, data["disclaimer"], size=7.6, color=COLORS["slate"])

    target = BytesIO()
    document.save(target)
    docx_bytes = target.getvalue()
    if output_path:
        with open(output_path, "wb") as output:
            output.write(docx_bytes)
    return docx_bytes
